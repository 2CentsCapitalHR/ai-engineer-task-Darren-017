import re
from pathlib import Path
from typing import List, Dict
from docx import Document
from core.rag import retrieve, ask_gemini

def analyze_document(path: Path, doc_type: str) -> List[Dict]:
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    issues: List[Dict] = []

    if re.search(r"UAE\s+Federal\s+Courts", text, re.I) and not re.search(r"ADGM\s+Courts", text, re.I):
        cites = retrieve("ADGM governing law jurisdiction clause for contracts")
        issues.append({
            "document": path.name,
            "section": "Jurisdiction",
            "issue": "Jurisdiction refers to UAE Federal Courts; expected ADGM Courts.",
            "severity": "High",
            "suggestion": "Use: 'This agreement is governed by ADGM law; disputes are subject to ADGM Courts.'",
            "citations": cites[:3],  
            "anchor_regex": r"UAE\s+Federal\s+Courts"
        })

    if re.search(r"\[[^\]]+\]", text):
        issues.append({
            "document": path.name,
            "section": "Placeholders",
            "issue": "Unresolved placeholders detected.",
            "severity": "Medium",
            "suggestion": "Replace placeholders like [Company], [Date], [Address] before submission.",
            "citations": [],
            "anchor_regex": r"\[[^\]]+\]"
        })

    if not re.search(r"Signature|Signed by|Authorised Signatory", text, re.I):
        cites = retrieve("ADGM execution signature requirements")
        issues.append({
            "document": path.name,
            "section": "Execution",
            "issue": "Signature/signatory section appears to be missing.",
            "severity": "High",
            "suggestion": "Add authorized signatory block: name, title, date (and seal if required).",
            "citations": cites[:2],
            "anchor_regex": r"$^"
        })

    if issues:
        bullet_points = "\n".join(f"- {i['section']}: {i['issue']} (Suggestion: {i['suggestion']})"
                                  for i in issues if i.get("issue"))
        phrased = ask_gemini(
            "Rewrite these compliance findings as succinct, professional review notes for a legal document.",
            bullet_points
        )
        if phrased and not phrased.startswith("(LLM error") and not phrased.startswith("(Gemini not configured)"):
            issues.append({
                "document": path.name,
                "section": "Summary",
                "issue": "LLM phrased review notes",
                "severity": "Info",
                "suggestion": phrased,
                "citations": []
            })
        else:
            pass

    return issues
